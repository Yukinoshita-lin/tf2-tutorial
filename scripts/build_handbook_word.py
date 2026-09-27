#!/usr/bin/env python3
"""Build the learning handbook DOCX (Word) and render the PDF with Microsoft Word.

Pipeline:
    1. pandoc:  docs/learning_handbook_zh.md -> docs/learning_handbook_zh.docx
                (GFM reader, TOC field, images embedded)
    2. python-docx post-pass:
                A4 page setup, CJK fonts (微软雅黑), code-block style,
                image down-scaling, table font size
    3. Word COM: update all fields (TOC page numbers) -> ExportAsFixedFormat PDF
                (falls back to LibreOffice soffice if Word is unavailable)

Usage:
    python scripts/build_handbook_word.py
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS = PROJECT_ROOT / "docs"
MD_PATH = DOCS / "learning_handbook_zh.md"
DOCX_PATH = DOCS / "learning_handbook_zh.docx"
PDF_PATH = DOCS / "learning_handbook_zh.pdf"

CJK_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
MAX_IMG_CM = 15.5  # A4 (21cm) 减去页边距后留一点余量


def _find_exe(name: str, candidates: list[str]) -> str:
    """在 PATH 与常见安装位置里找可执行文件（winget 装完后 PATH 可能未刷新）。"""
    import shutil
    found = shutil.which(name)
    if found:
        return found
    for c in candidates:
        if Path(c).exists():
            return c
    raise FileNotFoundError(f"{name} not found; tried PATH + {candidates}")


PANDOC = _find_exe("pandoc", [
    "C:/Users/32403/AppData/Local/Pandoc/pandoc.exe",
    "C:/Program Files/Pandoc/pandoc.exe",
])
SOFFICE = _find_exe("soffice", [
    "C:/Program Files/LibreOffice/program/soffice.exe",
    "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
])


def step1_pandoc() -> None:
    """markdown -> docx（在 docs/ 目录内运行，让 ../figures 相对路径可解析）。"""
    cmd = [
        PANDOC, MD_PATH.name,
        "-f", "gfm", "-t", "docx",
        "-o", DOCX_PATH.name,
        "--toc", "--toc-depth=2",
        "--standalone",
    ]
    print("[1/3] pandoc:", " ".join(cmd))
    proc = subprocess.run(cmd, cwd=DOCS, capture_output=True, text=True,
                          encoding="utf-8", errors="replace")
    if proc.returncode != 0:
        raise RuntimeError(f"pandoc failed:\n{proc.stdout}\n{proc.stderr}")
    print(f"      -> {DOCX_PATH} ({DOCX_PATH.stat().st_size / 1024:.0f} KB)")


def _set_style_font(style, ascii_font: str, cjk_font: str, size: float | None = None,
                    color=None) -> None:
    """设置样式的中西文字体（eastAsia 必须单独设，否则中文回落到默认字体）。"""
    try:
        font = style.font
    except Exception:
        return
    font.name = ascii_font
    if size is not None:
        font.size = size
    if color is not None:
        try:
            font.color.rgb = color
        except Exception:
            pass
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is None:
        rfonts = rpr.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts", {})
        rpr.append(rfonts)
    rfonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", cjk_font)


def step2_polish() -> None:
    """A4 页面、中文字体、代码块样式、图片缩放、表格字号。"""
    import docx  # python-docx
    from docx.shared import Cm, Pt, RGBColor
    from docx.enum.section import WD_ORIENT

    print("[2/3] python-docx 后处理...")
    doc = docx.Document(str(DOCX_PATH))

    # --- 页面: A4 ---
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin = sec.bottom_margin = Cm(2.0)

    # --- 样式字体 ---
    purple = RGBColor(0x5B, 0x21, 0xB6)
    styles = doc.styles
    normal = styles["Normal"]
    _set_style_font(normal, "Calibri", CJK_FONT, Pt(10.5))
    normal.paragraph_format.line_spacing = 1.3
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 15),
                       ("Heading 3", 12.5), ("Heading 4", 11)):
        if name in [s.name for s in styles]:
            st = styles[name]
            _set_style_font(st, "Calibri", CJK_FONT, Pt(size), purple)
    # 代码块样式（pandoc 的段落样式 Source Code / 字符样式 Verbatim Char）
    for name in ("Source Code", "Verbatim Char"):
        if name in [s.name for s in styles]:
            _set_style_font(styles[name], CODE_FONT, CJK_FONT, Pt(8.5))
    if "Block Text" in [s.name for s in styles]:  # 引用块
        st = styles["Block Text"]
        _set_style_font(st, "Calibri", CJK_FONT, Pt(10))
    if "Table" in [s.name for s in styles]:
        _set_style_font(styles["Table"], "Calibri", CJK_FONT, Pt(9))

    # --- 图片等比缩放到页宽内 ---
    max_w = Cm(MAX_IMG_CM)
    n_scaled = 0
    for shape in doc.inline_shapes:
        if shape.width > max_w:
            ratio = max_w / shape.width
            shape.height = int(shape.height * ratio)
            shape.width = int(max_w)
            n_scaled += 1
    # 表格内段落字号统一 9pt（pandoc 单元格有时不走 Table 样式）
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)
    doc.save(str(DOCX_PATH))
    print(f"      -> 缩放图片 {n_scaled} 张；表格 {len(doc.tables)} 个")
    print(f"      -> {DOCX_PATH} ({DOCX_PATH.stat().st_size / 1024:.0f} KB)")


def step3_word_pdf() -> None:
    """用 Microsoft Word 打开 → 更新目录字段 → 导出 PDF（用户指定的渲染引擎）。"""
    import win32com.client  # noqa: PLC0415

    print("[3/3] Microsoft Word 渲染 PDF...")
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(DOCX_PATH), ReadOnly=False)
        try:
            # 更新全部字段（含目录页码），再专门刷新目录
            doc.Fields.Update()
            for i in range(1, doc.TablesOfContents.Count + 1):
                doc.TablesOfContents(i).Update()
            doc.SaveAs2(str(DOCX_PATH))  # 保存更新后的目录
            doc.ExportAsFixedFormat(str(PDF_PATH), ExportFormat=17)  # wdExportFormatPDF
        finally:
            doc.Close(False)
    finally:
        word.Quit()
    print(f"      -> {PDF_PATH} ({PDF_PATH.stat().st_size / 1024:.0f} KB)")


def step3_libreoffice_fallback() -> None:
    """Word 不可用时的回退：LibreOffice headless 转换。"""
    print("[3/3] LibreOffice 渲染 PDF（Word 不可用，回退方案）...")
    subprocess.run(
        [SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", str(DOCS), str(DOCX_PATH)],
        check=True, capture_output=True, text=True, timeout=600)


def main() -> int:
    if not MD_PATH.exists():
        print(f"markdown 不存在: {MD_PATH}")
        return 2
    step1_pandoc()
    step2_polish()
    try:
        step3_word_pdf()
    except Exception as e:  # noqa: BLE001
        print(f"      Word COM 失败（{e}），尝试 LibreOffice 回退...")
        try:
            step3_libreoffice_fallback()
        except Exception as e2:  # noqa: BLE001
            print(f"      LibreOffice 也失败: {e2}")
            return 1
    print(f"\n完成:\n  DOCX {DOCX_PATH}\n  PDF  {PDF_PATH}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
